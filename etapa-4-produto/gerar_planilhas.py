"""Gera as planilhas de estratégia de produto da Calmaria Spas.

Saídas (na mesma pasta deste script):
- ai_project_canvas.csv   -> AI Project Canvas (Problema, Dados, IA, Experiência)
- matriz_csd.csv          -> Matriz CSD (Certezas, Suposições, Dúvidas)
- estrategia_produto.xlsx -> as duas abas num único arquivo (requer openpyxl)

Toda linha do canvas diz o TIPO da informação, para nunca misturar fato com aposta:
- Conhecido: verificado no enunciado ou no dataset (a fonte é citada)
- Hipótese: acreditamos, mas precisa ser validado
- Meta: valor-alvo proposto (não é resultado observado)
- A coletar: informação que ainda não existe
"""

import csv
from pathlib import Path

PASTA = Path(__file__).parent

COLUNAS_CANVAS = ["Pilar", "Bloco", "Conteúdo", "Tipo", "Fonte ou como obter"]

CANVAS = [
    # --- Problema -------------------------------------------------------------
    ["Problema", "Comportamento atual", "O assistente dá recomendações genéricas e respostas inconsistentes; não há visibilidade do comportamento em produção.", "Conhecido", "Enunciado do projeto"],
    ["Problema", "Evidências", "Dataset com 50 interações em 6 intenções e 6 categorias. As respostas de referência têm de 72 a 156 caracteres e citam tratamento, duração e preço.", "Conhecido", "dados/calmaria_interacoes.csv (validado em 2026-10-25)"],
    ["Problema", "Evidências", "Não existiam traces, scores ou métricas do assistente antes deste projeto.", "Conhecido", "Enunciado ('pouca visibilidade')"],
    ["Problema", "Evidências", "Score médio do juiz, % de casos críticos e acurácia de intenção do assistente atual.", "A coletar", "src/avaliar_respostas.py e src/avaliar_intencoes.py (Etapa 3)"],
    ["Problema", "Métricas", "Conversão do chat em agendamento; satisfação com o atendimento (CSAT); taxa de transbordo para atendente.", "A coletar", "Sistema de agendamento + pesquisa pós-atendimento"],
    ["Problema", "Métricas", "Score médio do juiz ≥ 4,0 e casos críticos ≤ 10% na amostra de avaliação.", "Meta", "Proposta para o Comitê; recalibrar após a primeira execução"],
    ["Problema", "Usuário-alvo", "Cliente que chega pelo canal digital e não sabe qual tratamento escolher, em especial o cliente novo (20 das 50 interações são de perfil 'novo').", "Hipótese", "Contagem é Conhecido (dataset); ser o público prioritário é Hipótese (ver proto-persona)"],
    # --- Dados ----------------------------------------------------------------
    ["Dados", "Disponíveis", "Pergunta do cliente, categoria, intenção, resposta de referência e perfil (novo, recorrente, vip).", "Conhecido", "Dataset"],
    ["Dados", "Disponíveis", "Traces do Langfuse (pergunta, resposta, versão do prompt, tokens, latência, custo) depois da Etapa 2.", "A coletar", "Execução do pipeline instrumentado"],
    ["Dados", "Ausentes", "Histórico de agendamentos por cliente; catálogo estruturado (preço, duração, contraindicações); agenda com disponibilidade; feedback explícito; conversões.", "A coletar", "Sistemas internos (a mapear com TI)"],
    ["Dados", "Qualidade", "Dataset pequeno, limpo e balanceado (6 a 10 exemplos por intenção, sem nulos nem duplicados). Perguntas bem escritas, provavelmente mais 'limpas' que o chat real.", "Conhecido", "Validação do dataset (tests/test_dataset.py)"],
    ["Dados", "Qualidade", "Os preços das respostas de referência são os vigentes.", "Hipótese", "Confirmar com o Comercial"],
    ["Dados", "Vieses", "Perfis desbalanceados (novo 20, recorrente 19, vip 11); benefício por perfil na referência (id 18); possível viés de gênero na linguagem (id 35); nenhum dado demográfico para medir equidade.", "Conhecido", "Dataset"],
    ["Dados", "Limitações", "50 linhas não sustentam treinar um modelo de recomendação; servem para avaliação. Respostas de referência afirmam ações e horários que o assistente não pode garantir.", "Conhecido", "Dataset + matriz de riscos (R04, R05)"],
    # --- IA -------------------------------------------------------------------
    ["IA", "Solução híbrida", "1) Regras de segurança (contraindicações e encaminhamento em casos de saúde); 2) recomendação por conteúdo (atributos do catálogo × necessidade extraída da pergunta); 3) filtragem colaborativa quando houver histórico; 4) LLM para entender o pedido e redigir a explicação.", "Hipótese", "Proposta de arquitetura; validar em shadow mode"],
    ["IA", "Entradas", "Pergunta, intenção classificada, categoria, perfil, histórico (quando houver) e catálogo oficial.", "Hipótese", "Depende de integração com o catálogo"],
    ["IA", "Saídas", "Até 2 tratamentos + motivo em uma frase + próximo passo; ou encaminhamento para atendente ou profissional.", "Meta", "Formato definido em prompts/assistente_improved.md"],
    ["IA", "Cold start", "Cliente novo sem histórico: popularidade por categoria + conteúdo. Tratamento novo sem avaliações: conteúdo + curadoria da coordenação técnica.", "Hipótese", "Validar com teste A/B"],
    ["IA", "Custo", "Custo por interação = tokens × preço do modelo, medido pelo Langfuse. Valor-alvo a definir com o volume diário real.", "A coletar", "Langfuse (Etapa 2) + volume de conversas"],
    ["IA", "Latência aceitável", "p95 ≤ 5 s por resposta no chat.", "Meta", "Proposta; validar com usuários e medir no Langfuse"],
    ["IA", "Riscos", "Alucinação de preço e horário, recomendação inadequada para condição de saúde, afirmar ação não executada, dependência de fornecedor.", "Conhecido", "etapa-1-governanca/matriz_riscos.csv (R03, R04, R05, R10)"],
    # --- Experiência ----------------------------------------------------------
    ["Experiência", "Pontos de contato", "Chat do site e do app; sugestões na página de cada tratamento; mensagem pós-sessão com sugestão de continuidade.", "Hipótese", "Validar com o time de Produto"],
    ["Experiência", "Feedback", "Botões 👍/👎 em cada recomendação, com motivo opcional ('não era isso', 'caro', 'já fiz'); nota pós-atendimento.", "Meta", "A implementar"],
    ["Experiência", "Transparência", "Aviso 'Você está falando com a assistente virtual da Calmaria'; link 'Por que esta sugestão?' com o motivo da recomendação.", "Meta", "Política de uso responsável §4"],
    ["Experiência", "Correção da recomendação", "Cliente pode responder 'não era isso' e receber alternativa; atendente pode corrigir e o caso vira exemplo de avaliação.", "Meta", "A implementar"],
    ["Experiência", "Fallback", "Técnico (API fora): mensagem de indisponibilidade + atendente. Qualidade (nota ≤ 2 na avaliação): regenerar com o prompt improved. Saúde: encaminhamento humano, sem recomendação.", "Conhecido", "Implementado no código da Etapa 3 (fallback por qualidade) e nos prompts; fallback técnico em produção a implementar"],
]

COLUNAS_CSD = ["Coluna", "Item", "Origem ou evidência", "Como validar ou próximo passo"]

CSD = [
    ["Certeza", "O assistente atual dá respostas genéricas e inconsistentes.", "Enunciado", "Quantificar com o score médio do juiz (Etapa 3)"],
    ["Certeza", "Não havia observabilidade do assistente em produção.", "Enunciado", "Resolvido pela instrumentação da Etapa 2"],
    ["Certeza", "Clientes mencionam informações de saúde no texto livre.", "Dataset: ids 21, 34, 43, 46", "Medir a frequência real nos traces (com cuidado de privacidade)"],
    ["Certeza", "As perguntas se dividem em 6 intenções e 6 categorias; 20 de 50 são de clientes novos.", "Dataset", "Comparar com a distribuição real do chat"],
    ["Certeza", "O assistente não tem integração com agenda nem com o catálogo.", "Arquitetura atual (só LLM)", "Mapear APIs disponíveis com TI"],
    ["Suposição", "Recomendações mais específicas aumentam a conversão em agendamento.", "Hipótese de produto", "Teste A/B: conversão por sessão com IA improved × atual"],
    ["Suposição", "Clientes novos são os que mais se beneficiam do assistente.", "Proto-persona", "Entrevistas + conversão por perfil"],
    ["Suposição", "O prompt improved gera respostas melhores que o baseline.", "Desenho do prompt", "Comparar no Langfuse e com o juiz (Etapas 2 e 3)"],
    ["Suposição", "O juiz LLM concorda com avaliadores humanos na maioria dos casos.", "Prática comum em LLM-as-Judge", "Calibração: humano dá nota às mesmas 10 respostas"],
    ["Suposição", "Clientes aceitam falar com uma IA se puderem pedir um humano.", "Hipótese de experiência", "Taxa de pedidos de atendente e CSAT no rollout"],
    ["Dúvida", "Qual a taxa de conversão atual do chat em agendamento?", "Não há dado", "Extrair do sistema de agendamento antes do rollout"],
    ["Dúvida", "Quantas conversas por dia o chat recebe?", "Não há dado", "Logs do canal; define o custo mensal"],
    ["Dúvida", "Existe catálogo estruturado com preços e contraindicações?", "Não há dado", "Levantamento com Comercial e Coordenação Técnica"],
    ["Dúvida", "Qual a base legal para tratar os dados de saúde mencionados?", "Matriz de riscos R02", "Parecer do Encarregado e do Jurídico"],
    ["Dúvida", "Qual latência o cliente tolera no chat?", "Meta de 5 s é proposta", "Teste com usuários + abandono por faixa de latência"],
    ["Dúvida", "Os preços das respostas de referência são os vigentes?", "Dataset", "Confirmar com o Comercial"],
]


def salvar_csv(nome, colunas, linhas):
    with open(PASTA / nome, "w", encoding="utf-8", newline="") as f:
        escritor = csv.writer(f)
        escritor.writerow(colunas)
        escritor.writerows(linhas)


CORES_TIPO = {"Conhecido": "C6EFCE", "Hipótese": "FFEB9C", "Meta": "BDD7EE", "A coletar": "F8CBAD",
              "Certeza": "C6EFCE", "Suposição": "FFEB9C", "Dúvida": "F8CBAD"}


def salvar_xlsx(nome):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    abas = [("AI Project Canvas", COLUNAS_CANVAS, CANVAS, 3), ("Matriz CSD", COLUNAS_CSD, CSD, 0)]
    for i, (titulo, colunas, linhas, coluna_cor) in enumerate(abas):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = titulo
        ws.append(colunas)
        for linha in linhas:
            ws.append(linha)
        for celula in ws[1]:
            celula.font = Font(bold=True)
        for coluna in ws.columns:
            largura = 70 if coluna[0].value in ("Conteúdo", "Item") else 28
            ws.column_dimensions[coluna[0].column_letter].width = largura
            for celula in coluna:
                celula.alignment = Alignment(wrap_text=True, vertical="top")
        for linha in ws.iter_rows(min_row=2):
            cor = CORES_TIPO.get(linha[coluna_cor].value)
            if cor:
                linha[coluna_cor].fill = PatternFill("solid", fgColor=cor)
        ws.freeze_panes = "A2"
    wb.save(PASTA / nome)


if __name__ == "__main__":
    salvar_csv("ai_project_canvas.csv", COLUNAS_CANVAS, CANVAS)
    salvar_csv("matriz_csd.csv", COLUNAS_CSD, CSD)
    try:
        salvar_xlsx("estrategia_produto.xlsx")
    except ImportError:
        print("openpyxl não instalado: .xlsx não gerado (os CSVs foram gerados).")
    print("Planilhas geradas em", PASTA)
    from collections import Counter
    print("Canvas por pilar:", dict(Counter(l[0] for l in CANVAS)))
    print("Canvas por tipo: ", dict(Counter(l[3] for l in CANVAS)))
    print("CSD por coluna:  ", dict(Counter(l[0] for l in CSD)))
